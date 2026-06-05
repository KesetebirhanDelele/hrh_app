from __future__ import annotations

import json
from pathlib import Path
from typing import Any, Dict, List, Optional

from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH


def _read_json(path: Path) -> Dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8"))


def _add_heading(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def _add_paragraph(doc: Document, text: str, bold_prefix: Optional[str] = None) -> None:
    p = doc.add_paragraph()
    if bold_prefix:
        run = p.add_run(bold_prefix)
        run.bold = True
    p.add_run(text)


def _add_citations(doc: Document, citations: List[Dict[str, Any]]) -> None:
    if not citations:
        return
    doc.add_paragraph("Citations:", style=None).runs[0].bold = True
    for c in citations:
        title = c.get("source_title", "Unknown source")
        locator = c.get("locator", "")
        url = c.get("source_url")
        published = c.get("published_date")
        parts = [title]
        if published:
            parts.append(f"({published})")
        if locator:
            parts.append(f"— {locator}")
        if url:
            parts.append(f"— {url}")
        doc.add_paragraph(" ".join(parts), style="List Bullet")
        quote = c.get("quote")
        if quote:
            doc.add_paragraph(f'Quote: "{quote}"', style="List Bullet")


def render_docx(job_id: str, payload: Dict[str, Any], out_path: Path) -> Path:
    doc = Document()

    # Basic styling
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    title = job_id
    doc.add_heading(title, level=0)

    # Meta
    meta_lines = []
    if payload.get("spec_id"):
        meta_lines.append(f"spec_id: {payload['spec_id']}")
    if payload.get("generated_at"):
        meta_lines.append(f"generated_at: {payload['generated_at']}")
    if payload.get("country", {}).get("name"):
        c = payload["country"]
        meta_lines.append(f"country: {c.get('name')}{' (' + c.get('iso3') + ')' if c.get('iso3') else ''}")
    if meta_lines:
        p = doc.add_paragraph("\n".join(meta_lines))
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT

    doc.add_paragraph("")

    if job_id == "phase1_discovery_qa":
        _add_heading(doc, "Discovery questions and answers", level=1)
        for q in payload.get("questions", []):
            _add_heading(doc, f"{q.get('question_id','')}", level=2)
            _add_paragraph(doc, q.get("question", ""), bold_prefix="Q: ")
            _add_paragraph(doc, q.get("answer", ""), bold_prefix="A: ")

            ev = q.get("evidence", {})
            _add_paragraph(doc, str(ev.get("quality", "")), bold_prefix="Evidence quality: ")
            _add_paragraph(doc, ev.get("rationale", ""), bold_prefix="Rationale: ")
            _add_citations(doc, ev.get("citations", []) or [])
            doc.add_paragraph("")

    elif job_id == "country_learning_briefs":
        _add_heading(doc, "Learning domain briefs", level=1)
        for d in payload.get("domains", []):
            _add_heading(doc, f"{d.get('domain_id','')}: {d.get('domain','')}", level=2)
            _add_paragraph(doc, d.get("summary", ""), bold_prefix="Summary: ")
            doc.add_paragraph("Key questions:", style=None).runs[0].bold = True
            for kq in d.get("key_questions", []) or []:
                doc.add_paragraph(str(kq), style="List Bullet")

            ev = d.get("evidence", {})
            _add_paragraph(doc, str(ev.get("quality", "")), bold_prefix="Evidence quality: ")
            _add_paragraph(doc, ev.get("rationale", ""), bold_prefix="Rationale: ")
            _add_citations(doc, ev.get("citations", []) or [])
            doc.add_paragraph("")

    elif job_id == "rrr_evidence_matrix":
        _add_heading(doc, "RRR Solutions — Evidence Matrix", level=1)
        for sol in payload.get("solutions", []):
            label = f"{sol.get('solution_id', '')}: {sol.get('solution', '')}"
            _add_heading(doc, label, level=2)
            _add_paragraph(doc, sol.get("mechanism", ""), bold_prefix="Mechanism: ")

            feas = sol.get("feasibility_resource_constrained", "")
            if feas:
                _add_paragraph(doc, feas, bold_prefix="Feasibility (resource-constrained): ")

            risks = sol.get("risks", "")
            if risks:
                _add_paragraph(doc, risks, bold_prefix="Risks: ")

            notes = sol.get("implementation_notes", "")
            if notes:
                _add_paragraph(doc, notes, bold_prefix="Implementation notes: ")

            ev = sol.get("evidence", {})
            if ev:
                _add_paragraph(doc, str(ev.get("quality", "")), bold_prefix="Evidence quality: ")
                _add_paragraph(doc, ev.get("rationale", ""), bold_prefix="Rationale: ")
                _add_citations(doc, ev.get("citations", []) or [])

            doc.add_paragraph("")

    else:
        raise KeyError(f"render_docx: unsupported job_id '{job_id}'")

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)
    return out_path


def render_docx_file(job_id: str, input_json_path: str, output_docx_path: Optional[str] = None) -> Path:
    inp = Path(input_json_path).resolve()
    if not inp.exists():
        raise FileNotFoundError(f"Input file not found: {inp}")
    payload = _read_json(inp)
    out = Path(output_docx_path).resolve() if output_docx_path else inp.with_suffix(".docx")
    return render_docx(job_id, payload, out)
