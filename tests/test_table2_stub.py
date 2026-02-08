from app.jobs.table2_stub_runner import run_table2_stub
from app.core.validators import validate_output
from app.jobs.registry import get_job
from app.render.xlsx import render_xlsx
import json
from pathlib import Path
from openpyxl import load_workbook
import tempfile


def test_table2_stub_runs_and_validates() -> None:
    out_path = run_table2_stub(spec_id="table2_root_cause_framework_v1")
    job = get_job("table2_root_cause_mapping")
    payload = json.loads(out_path.read_text(encoding="utf-8"))
    schema_rel = str(job.output_schema.relative_to(Path.cwd()))
    validate_output(payload, schema_rel)


def test_table2_xlsx_renders_from_payload() -> None:
    """Regression test: ensure Table2 XLSX renders actual payload data, not template/placeholders."""
    # Create a minimal test payload with specific, verifiable data
    test_payload = {
        "job_id": "table2_root_cause_mapping",
        "spec_id": "test_spec",
        "generated_at": "2024-01-01",
        "framework_items": [
            {
                "item_id": "RC-TEST-001",
                "category": "Test Category",
                "root_cause": "Test Root Cause",
                "definition": "This is a very specific test definition that should appear in the XLSX",
                "evidence": {
                    "quality": "high",
                    "rationale": "Test rationale for verification",
                    "citations": [
                        {
                            "source_title": "Test Source",
                            "locator": "Page 42",
                            "source_url": "https://example.com/test"
                        }
                    ]
                },
                "tags": ["test_tag_1", "test_tag_2"]
            }
        ]
    }

    # Render to XLSX
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
        tmp_path = Path(tmp.name)

    try:
        render_xlsx("table2_root_cause_mapping", test_payload, tmp_path)

        # Load and verify the XLSX contains the payload data
        wb = load_workbook(tmp_path)
        ws = wb.active

        # Check headers (row 1)
        assert ws.cell(1, 1).value == "item_id"
        assert ws.cell(1, 4).value == "definition"
        assert ws.cell(1, 5).value == "evidence_quality"

        # Check first data row (row 2) - this is the critical test
        assert ws.cell(2, 1).value == "RC-TEST-001", "item_id should match payload"
        assert ws.cell(2, 2).value == "Test Category", "category should match payload"
        assert ws.cell(2, 3).value == "Test Root Cause", "root_cause should match payload"
        assert ws.cell(2, 4).value == "This is a very specific test definition that should appear in the XLSX", \
            "definition should match payload exactly, not be a placeholder"
        assert ws.cell(2, 5).value == "high", "evidence quality should match payload"
        assert ws.cell(2, 6).value == "Test rationale for verification", "evidence rationale should match payload"
        assert "Test Source" in ws.cell(2, 7).value, "citations should include source title"
        assert "Page 42" in ws.cell(2, 7).value, "citations should include locator"
        assert ws.cell(2, 8).value == "test_tag_1; test_tag_2", "tags should be semicolon-separated"

        wb.close()
    finally:
        # Clean up
        if tmp_path.exists():
            tmp_path.unlink()
